"""Extracción de texto de inventarios Word (.docx; .doc solo metadatos/aviso)."""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

INVENTARIO_RE = re.compile(r"inventario", re.IGNORECASE)


@dataclass
class InventarioDoc:
    path: str
    fondo_guess: str
    format: str
    paragraphs: list[str] = field(default_factory=list)
    tables_as_text: list[str] = field(default_factory=list)
    error: str | None = None

    @property
    def text(self) -> str:
        parts = self.paragraphs + self.tables_as_text
        return "\n".join(p for p in parts if p and p.strip())


def find_inventario_word_files(root: Path) -> list[Path]:
    hits: list[Path] = []
    for p in root.rglob("*"):
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".doc", ".docx"}:
            continue
        if INVENTARIO_RE.search(p.name):
            hits.append(p)
    return sorted(hits)


def _fondo_guess(path: Path, root: Path) -> str:
    try:
        rel = path.relative_to(root)
        parts = rel.parts
        # .../Fondos/<Fondo>/...
        for i, part in enumerate(parts):
            if part.lower() == "fondos" and i + 1 < len(parts):
                return parts[i + 1]
        return parts[0] if parts else path.parent.name
    except ValueError:
        return path.parent.name


def read_docx(path: Path) -> InventarioDoc:
    try:
        from docx import Document  # type: ignore
    except ImportError as e:
        return InventarioDoc(
            path=str(path),
            fondo_guess=path.parent.name,
            format="docx",
            error=f"python-docx no instalado: {e}",
        )

    try:
        doc = Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        tables: list[str] = []
        for ti, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [" ".join(c.text.split()) for c in row.cells]
                rows.append(" | ".join(cells))
            tables.append(f"[tabla {ti+1}]\n" + "\n".join(rows))
        return InventarioDoc(
            path=str(path),
            fondo_guess=path.parent.name,
            format="docx",
            paragraphs=paragraphs,
            tables_as_text=tables,
        )
    except Exception as e:  # noqa: BLE001
        return InventarioDoc(
            path=str(path),
            fondo_guess=path.parent.name,
            format="docx",
            error=str(e),
        )


def read_doc_legacy(path: Path) -> InventarioDoc:
    """Los .doc binarios no se parsean aquí; se registra la ruta para conversión local."""
    return InventarioDoc(
        path=str(path),
        fondo_guess=path.parent.name,
        format="doc",
        error=(
            "Formato .doc (Word 97–2003) no parseado en este entorno. "
            "Convertir a .docx (LibreOffice: soffice --headless --convert-to docx) y reintentar."
        ),
    )


def extract_inventarios(root: Path) -> list[InventarioDoc]:
    docs: list[InventarioDoc] = []
    for path in find_inventario_word_files(root):
        fondo = _fondo_guess(path, root)
        if path.suffix.lower() == ".docx":
            inv = read_docx(path)
        else:
            inv = read_doc_legacy(path)
        inv.fondo_guess = fondo
        docs.append(inv)
    return docs


def write_inventario_outputs(docs: list[InventarioDoc], out_dir: Path) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    index: list[dict[str, Any]] = []
    paths: dict[str, Path] = {}
    for doc in docs:
        safe = re.sub(r"[^\w\-]+", "_", doc.fondo_guess).strip("_") or "fondo"
        base = out_dir / f"inventario__{safe}"
        # Evitar colisiones
        candidate = Path(str(base) + ".md")
        n = 2
        while candidate.exists():
            candidate = Path(str(base) + f"_{n}.md")
            n += 1

        md_lines = [
            f"# Inventario — {doc.fondo_guess}",
            "",
            f"- Archivo: `{doc.path}`",
            f"- Formato: `{doc.format}`",
        ]
        if doc.error:
            md_lines += ["", f"**Error/aviso:** {doc.error}", ""]
        else:
            md_lines += ["", "## Contenido extraído", "", doc.text or "_vacío_", ""]
        candidate.write_text("\n".join(md_lines), encoding="utf-8")
        paths[doc.path] = candidate
        index.append(
            {
                "fondo_guess": doc.fondo_guess,
                "source": doc.path,
                "format": doc.format,
                "extracted_md": str(candidate),
                "error": doc.error,
                "n_paragraphs": len(doc.paragraphs),
                "n_tables": len(doc.tables_as_text),
            }
        )

    index_path = out_dir / "inventarios_index.json"
    index_path.write_text(json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    paths["__index__"] = index_path
    return paths
